import io
import json
import re
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ILLEGAL_XML_CHARS_RE = re.compile(r'[\x00-\x08\x0b-\x0c\x0e-\x1f\x7f-\x84\x86-\x9f]')


def sanitize_cell_value(val):
    """
    Sanitizes values for OpenXML / openpyxl cell placement.
    Returns None for empty/null values to avoid writing invalid <c t="inlineStr"/> tags.
    """
    if val is None:
        return None
    if isinstance(val, bool):
        return val
    if isinstance(val, (int, float)):
        return val
    if isinstance(val, (dict, list)):
        val = json.dumps(val, ensure_ascii=False)
    
    val_str = str(val)
    val_str = ILLEGAL_XML_CHARS_RE.sub('', val_str)
    
    if not val_str:
        return None
    return val_str


class ExcelExportService:
    @staticmethod
    def generate_excel_bytes(records: list) -> bytes:
        """
        Formats history records into an in-memory Excel workbook (.xlsx)
        and returns the raw bytes.
        """
        wb = openpyxl.Workbook()
        
        # ── Sheet 1: Post History Data ──────────────────────────────────────────
        ws = wb.active
        ws.title = "Post History"

        columns = [
            ("id",              "History ID",        22),
            ("created_at",      "Created At",       20),
            ("topic",           "Topic",            32),
            ("post_type",       "Post Type",        16),
            ("tone",            "Tone",             16),
            ("language",        "Language",         14),
            ("length",          "Length",           14),
            ("target_audience", "Target Audience",  24),
            ("writing_style",   "Writing Style",    28),
            ("personal_context","Personal Context", 30),
            ("key_points",      "Key Points",       30),
            ("post",            "Post Content",     75),
            ("word_count",      "Word Count",       12),
            ("reading_time",    "Reading Time",     15),
            ("emoji_count",     "Emoji Count",      13),
            ("hashtag_count",   "Hashtag Count",    14),
            ("action",          "Action",           15),
            ("parent_id",       "Parent ID",        22),
        ]

        # Style elements
        header_fill = PatternFill("solid", fgColor="0B66C2")  # WriteFlow / LinkedIn Blue
        header_font = Font(color="FFFFFF", bold=True, name="Segoe UI", size=10)
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        header_border = Border(
            left=Side(style="thin", color="D0D7DE"),
            right=Side(style="thin", color="D0D7DE"),
            bottom=Side(style="medium", color="0B66C2")
        )

        # Write Headers
        for col_idx, (_, label, width) in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=col_idx, value=label)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_align
            cell.border = header_border
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        ws.row_dimensions[1].height = 28
        ws.freeze_panes = "A2"

        # Write Data Rows
        wrap_alignment = Alignment(vertical="top", wrap_text=True)
        nowrap_alignment = Alignment(vertical="top", wrap_text=False)
        even_row_fill = PatternFill("solid", fgColor="F7F9FC")

        rows_written = 0
        for row_idx, rec in enumerate(records, start=2):
            row_fill = even_row_fill if row_idx % 2 == 0 else None
            
            for col_idx, (field_name, _, _) in enumerate(columns, start=1):
                raw_val = rec.get(field_name)
                val = sanitize_cell_value(raw_val)
                
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                
                should_wrap = field_name in ("post", "topic", "personal_context", "key_points", "writing_style")
                cell.alignment = wrap_alignment if should_wrap else nowrap_alignment
                if row_fill:
                    cell.fill = row_fill

            post_str = str(rec.get("post", "") or "")
            newlines = post_str.count("\n") + 1
            calc_height = max(20, min(newlines * 16, 150))
            ws.row_dimensions[row_idx].height = calc_height
            rows_written += 1

        last_col_letter = get_column_letter(len(columns))
        ws.auto_filter.ref = f"A1:{last_col_letter}{max(rows_written + 1, 2)}"

        # ── Sheet 2: Export Summary ─────────────────────────────────────────────
        ws_summary = wb.create_sheet("Export Summary")
        summary_bold = Font(bold=True, name="Segoe UI", size=11)
        
        summary_data = [
            ("Export Execution Date", datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
            ("Total Records Exported", len(records)),
            ("Total Rows Written", rows_written),
            ("Full Post Content Included", "YES (Un-truncated in 'Post Content' column)"),
            ("Database Source", "WriteFlow History Database"),
        ]

        for r_i, (k, v) in enumerate(summary_data, start=1):
            ws_summary.cell(row=r_i, column=1, value=k).font = summary_bold
            ws_summary.cell(row=r_i, column=2, value=v)

        ws_summary.column_dimensions["A"].width = 28
        ws_summary.column_dimensions["B"].width = 65

        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        return stream.getvalue()
