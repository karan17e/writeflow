"""
WriteFlow Production History Standalone Export Utility
======================================================
Fetches live history data directly from the WriteFlow production API:
  GET https://writeflow-yuwk.onrender.com/api/history

Exports all available records to an Excel (.xlsx) file with full,
un-truncated post content and formatted columns.

Usage:
  python export_history.py
  -- or --
  ../venv/Scripts/python.exe export_history.py
"""

import sys
import json
import datetime
import urllib.request
import urllib.error
import os

PRODUCTION_BASE_URL = "https://writeflow-yuwk.onrender.com"
HISTORY_ENDPOINT = f"{PRODUCTION_BASE_URL}/api/history"
PAGE_LIMIT = 100

timestamp_str = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
OUTPUT_FILENAME = f"writeflow_history_export_{timestamp_str}.xlsx"


def fetch_all_production_history() -> list:
    """
    Fetch all history records from production API, handling pagination.
    Uses a 120s timeout to allow for Render Free tier cold start wakeups.
    """
    all_records = []
    skip = 0

    print("=" * 65)
    print("  WriteFlow Production History Standalone Exporter")
    print("=" * 65)
    print(f"[1/3] Connecting to Production API:\n      {HISTORY_ENDPOINT}")

    while True:
        url = f"{HISTORY_ENDPOINT}?limit={PAGE_LIMIT}&skip={skip}"
        print(f"      -> Fetching page (skip={skip}, limit={PAGE_LIMIT})...")

        try:
            req = urllib.request.Request(
                url,
                headers={
                    "Accept": "application/json",
                    "User-Agent": "WriteFlow-HistoryExporter/1.0"
                }
            )
            # 120 second timeout for Render Free cold starts
            with urllib.request.urlopen(req, timeout=120) as response:
                if response.status != 200:
                    print(f"      ✗ Received non-200 status code: {response.status}")
                    break

                raw = response.read().decode("utf-8")
                records = json.loads(raw)

                if not isinstance(records, list):
                    print("      ✗ Unexpected response format (expected a JSON array).")
                    break

                if not records:
                    # Empty list returned, end of pagination
                    break

                all_records.extend(records)
                print(f"      ✓ Retrieved {len(records)} records (Total so far: {len(all_records)})")

                if len(records) < PAGE_LIMIT:
                    # Last page reached
                    break

                skip += PAGE_LIMIT

        except urllib.error.HTTPError as e:
            print(f"      ✗ HTTP Error {e.code}: {e.reason}")
            sys.exit(1)
        except urllib.error.URLError as e:
            print(f"      ✗ Network Error: {e.reason}")
            print("      Please ensure internet connectivity and that WriteFlow service is up.")
            sys.exit(1)
        except json.JSONDecodeError as e:
            print(f"      ✗ Failed to parse JSON response: {e}")
            sys.exit(1)

    print(f"      ✓ Fetch complete. Total history records retrieved: {len(all_records)}")
    return all_records


def ensure_openpyxl():
    """Ensure openpyxl is installed and return the module."""
    try:
        import openpyxl
        return openpyxl
    except ImportError:
        print("\n[ERROR] 'openpyxl' package is missing.")
        print("Please install it using:")
        print("  pip install openpyxl")
        print("  -- or --")
        print("  ..\\venv\\Scripts\\pip.exe install openpyxl")
        sys.exit(1)


def generate_excel_export(records: list, openpyxl_mod) -> str:
    """
    Format records into an Excel sheet with headers, wrap text, 
    custom column widths, and a summary sheet.
    """
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl_mod.Workbook()
    
    # ── Sheet 1: Post History Data ──────────────────────────────────────────
    ws = wb.active
    ws.title = "Post History"

    columns = [
        ("id",              "History ID",        22),
        ("created_at",      "Created At",       20),
        ("topic",           "Topic",            32),
        ("post_type",       "Post Type",        16),
        ("tone",            "Tone",             16),
        ("language",        "Language",         14),
        ("length",          "Length",           14),
        ("target_audience", "Target Audience",  24),
        ("writing_style",   "Writing Style",    28),
        ("personal_context","Personal Context", 30),
        ("key_points",      "Key Points",       30),
        ("post",            "Post Content",     75),
        ("word_count",      "Word Count",       12),
        ("reading_time",    "Reading Time",     15),
        ("emoji_count",     "Emoji Count",      13),
        ("hashtag_count",   "Hashtag Count",    14),
        ("action",          "Action",           15),
        ("parent_id",       "Parent ID",        22),
    ]

    # Style elements
    header_fill = PatternFill("solid", fgColor="0B66C2")  # WriteFlow / LinkedIn Blue
    header_font = Font(color="FFFFFF", bold=True, name="Segoe UI", size=10)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    header_border = Border(
        left=Side(style="thin", color="D0D7DE"),
        right=Side(style="thin", color="D0D7DE"),
        bottom=Side(style="medium", color="0B66C2")
    )

    # Write Headers
    for col_idx, (_, label, width) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = header_border
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"

    # Write Data
    wrap_alignment = Alignment(vertical="top", wrap_text=True)
    nowrap_alignment = Alignment(vertical="top", wrap_text=False)
    even_row_fill = PatternFill("solid", fgColor="F7F9FC")

    for row_idx, rec in enumerate(records, start=2):
        row_fill = even_row_fill if row_idx % 2 == 0 else None
        
        for col_idx, (field_name, _, _) in enumerate(columns, start=1):
            val = rec.get(field_name, "")
            if val is None:
                val = ""
            
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            
            # Text wrapping for long text fields
            should_wrap = field_name in ("post", "topic", "personal_context", "key_points", "writing_style")
            cell.alignment = wrap_alignment if should_wrap else nowrap_alignment
            if row_fill:
                cell.fill = row_fill

        # Adjust row height based on post content length
        post_str = str(rec.get("post", "") or "")
        newlines = post_str.count("\n") + 1
        calc_height = max(20, min(newlines * 16, 150))
        ws.row_dimensions[row_idx].height = calc_height

    # Auto-filter on all data columns
    last_col_letter = get_column_letter(len(columns))
    ws.auto_filter.ref = f"A1:{last_col_letter}{max(len(records) + 1, 2)}"

    # ── Sheet 2: Export Summary ─────────────────────────────────────────────
    ws_summary = wb.create_sheet("Export Summary")
    summary_bold = Font(bold=True, name="Segoe UI", size=11)
    
    summary_data = [
        ("Export Execution Date", datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Production API Base URL", PRODUCTION_BASE_URL),
        ("History Endpoint", HISTORY_ENDPOINT),
        ("Total Records Exported", len(records)),
        ("Output File Name", OUTPUT_FILENAME),
        ("Full Post Content Included", "YES (Un-truncated in 'Post Content' column)"),
        ("Database Source", "Live WriteFlow Production API"),
    ]

    for r_i, (k, v) in enumerate(summary_data, start=1):
        ws_summary.cell(row=r_i, column=1, value=k).font = summary_bold
        ws_summary.cell(row=r_i, column=2, value=v)

    ws_summary.column_dimensions["A"].width = 28
    ws_summary.column_dimensions["B"].width = 65

    # Save output file
    output_path = os.path.abspath(OUTPUT_FILENAME)
    wb.save(output_path)
    return output_path


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")

    records = fetch_all_production_history()

    print("\n[2/3] Checking dependencies...")
    openpyxl_mod = ensure_openpyxl()
    print(f"      ✓ openpyxl {openpyxl_mod.__version__} is ready.")

    print("\n[3/3] Exporting to Excel format...")
    file_path = generate_excel_export(records, openpyxl_mod)
    
    print("\n" + "=" * 65)
    print("  EXPORT COMPLETED SUCCESSFULLY")
    print("=" * 65)
    print(f"  Output File     : {OUTPUT_FILENAME}")
    print(f"  Full Path       : {file_path}")
    print(f"  Records Exported: {len(records)}")
    print(f"  Source Endpoint : {HISTORY_ENDPOINT}")
    print("=" * 65)


if __name__ == "__main__":
    main()
