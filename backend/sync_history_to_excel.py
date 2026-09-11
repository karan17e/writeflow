import urllib.request
import urllib.error
import json
import time
import os
import re
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

API_URL = "https://writeflow-yuwk.onrender.com/api/history"
EXCEL_FILE = r"E:\karan ka linkdin\backend\WriteFlow_History.xlsx"
POLL_INTERVAL_SEC = 30

ILLEGAL_XML_CHARS_RE = re.compile(r'[\x00-\x08\x0b-\x0c\x0e-\x1f\x7f-\x84\x86-\x9f]')

def sanitize_cell_value(val):
    if val is None: return None
    if isinstance(val, bool): return val
    if isinstance(val, (int, float)): return val
    if isinstance(val, (dict, list)): val = json.dumps(val, ensure_ascii=False)
    val_str = str(val)
    val_str = ILLEGAL_XML_CHARS_RE.sub('', val_str)
    return val_str if val_str else None

COLUMNS = [
    ("id",              "History ID",        22),
    ("created_at",      "Created At",       20),
    ("topic",           "Topic",            32),
    ("post_type",       "Post Type",        16),
    ("tone",            "Tone",             16),
    ("language",        "Language",         14),
    ("length",          "Length",           14),
    ("target_audience", "Target Audience",  24),
    ("writing_style",   "Writing Style",    28),
    ("personal_context","Personal Context", 30),
    ("key_points",      "Key Points",       30),
    ("post",            "Post Content",     75),
    ("word_count",      "Word Count",       12),
    ("reading_time",    "Reading Time",     15),
    ("emoji_count",     "Emoji Count",      13),
    ("hashtag_count",   "Hashtag Count",    14),
    ("action",          "Action",           15),
    ("parent_id",       "Parent ID",        22),
]

def apply_headers_and_styles(ws):
    header_fill = PatternFill("solid", fgColor="0B66C2")
    header_font = Font(color="FFFFFF", bold=True, name="Segoe UI", size=10)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    header_border = Border(
        left=Side(style="thin", color="D0D7DE"),
        right=Side(style="thin", color="D0D7DE"),
        bottom=Side(style="medium", color="0B66C2")
    )
    for col_idx, (_, label, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = header_border
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"

def write_row(ws, row_idx, rec):
    wrap_alignment = Alignment(vertical="top", wrap_text=True)
    nowrap_alignment = Alignment(vertical="top", wrap_text=False)
    even_row_fill = PatternFill("solid", fgColor="F7F9FC")
    row_fill = even_row_fill if row_idx % 2 == 0 else None
    
    for col_idx, (field_name, _, _) in enumerate(COLUMNS, start=1):
        raw_val = rec.get(field_name)
        val = sanitize_cell_value(raw_val)
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        should_wrap = field_name in ("post", "topic", "personal_context", "key_points", "writing_style")
        cell.alignment = wrap_alignment if should_wrap else nowrap_alignment
        if row_fill:
            cell.fill = row_fill

    post_str = str(rec.get("post", "") or "")
    newlines = post_str.count("\n") + 1
    calc_height = max(20, min(newlines * 16, 150))
    ws.row_dimensions[row_idx].height = calc_height

def are_records_equal(remote_rec, excel_rec):
    for key, _, _ in COLUMNS:
        remote_val = sanitize_cell_value(remote_rec.get(key))
        excel_val = excel_rec.get(key)
        
        # Handle cases where excel reads empty strings as None or vice versa
        if remote_val == "" and excel_val is None:
            continue
        if excel_val == "" and remote_val is None:
            continue
            
        if remote_val != excel_val:
            if str(remote_val) != str(excel_val):
                return False
    return True

def fetch_all_history():
    records = []
    skip = 0
    limit = 100
    while True:
        url = f"{API_URL}?skip={skip}&limit={limit}"
        req = urllib.request.Request(url)
        try:
            with urllib.request.urlopen(req) as response:
                if response.status == 200:
                    data = json.loads(response.read().decode('utf-8'))
                    if not data:
                        break
                    records.extend(data)
                    skip += limit
                else:
                    print(f"[ERROR] API returned status {response.status}")
                    break
        except urllib.error.URLError as e:
            print(f"[ERROR] Connection failed: {e}")
            return None # Return None on failure so we don't assume 0 records
    return records

def sync():
    print("[SYNC] Checking production History...")
    remote_records = fetch_all_history()
    
    if remote_records is None:
        print("[SYNC] Failed to fetch history. Retrying next interval.")
        return
        
    print(f"[SYNC] {len(remote_records)} records found.")
    
    remote_dict = {str(rec['id']): rec for rec in remote_records}
    
    # Load or Create Excel
    if os.path.exists(EXCEL_FILE):
        try:
            wb = openpyxl.load_workbook(EXCEL_FILE)
            ws = wb.active
        except Exception as e:
            print(f"[ERROR] Could not read Excel file: {e}")
            return
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Post History"
        apply_headers_and_styles(ws)
        
    # Map existing rows by ID
    existing_rows = {}
    existing_records = {}
    
    headers = [col[0] for col in COLUMNS]
    max_row = 1
    
    for row_idx, row_cells in enumerate(ws.iter_rows(min_row=2, max_col=len(COLUMNS), values_only=True), start=2):
        cell_val = row_cells[0]
        if cell_val:
            str_id = str(cell_val)
            existing_rows[str_id] = row_idx
            
            record_dict = {}
            for col_idx, key in enumerate(headers):
                val = row_cells[col_idx]
                record_dict[key] = val
            existing_records[str_id] = record_dict
            
        if row_idx > max_row:
            max_row = row_idx

    next_row_idx = max(existing_rows.values()) + 1 if existing_rows else 2
    
    changes_detected = False
    new_records = 0
    updated_records = 0
    
    for remote_id, remote_rec in remote_dict.items():
        if remote_id not in existing_rows:
            # New record
            write_row(ws, next_row_idx, remote_rec)
            existing_rows[remote_id] = next_row_idx
            next_row_idx += 1
            new_records += 1
            changes_detected = True
        else:
            # Compare existing record
            existing_rec = existing_records.get(remote_id, {})
            if not are_records_equal(remote_rec, existing_rec):
                row_idx = existing_rows[remote_id]
                write_row(ws, row_idx, remote_rec)
                updated_records += 1
                changes_detected = True
                
    if changes_detected:
        try:
             last_col_letter = get_column_letter(len(COLUMNS))
             ws.auto_filter.ref = f"A1:{last_col_letter}{max(ws.max_row, 2)}"
             wb.save(EXCEL_FILE)
             
             parts = []
             if new_records > 0:
                 parts.append(f"{new_records} new record(s)")
             if updated_records > 0:
                 parts.append(f"{updated_records} updated record(s)")
                 
             print(f"[SYNC] {' and '.join(parts)} detected.")
             print(f"[SYNC] Excel updated: {EXCEL_FILE}")
        except Exception as e:
             print(f"[ERROR] Failed to save Excel file: {e}")
    else:
        print("[SYNC] No changes detected.")

if __name__ == "__main__":
    print(f"Starting Background Sync to: {EXCEL_FILE}")
    print(f"Target API: {API_URL}")
    print(f"Polling Interval: {POLL_INTERVAL_SEC} seconds")
    print("Press Ctrl+C to stop.\n")
    try:
        # Run immediately on start
        sync()
        while True:
            time.sleep(POLL_INTERVAL_SEC)
            sync()
    except KeyboardInterrupt:
        print("\n[SYNC] Stopped by user.")
